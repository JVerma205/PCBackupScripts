import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT_DIR = Path(__file__).resolve().parent
LOGS_DIR = ROOT_DIR / "logs"
OUTPUT_FILE = ROOT_DIR / "arch_deletion_stats.xlsx"


def iter_log_csv_files(logs_dir: Path):
    if not logs_dir.exists():
        return []

    files = []
    for path in sorted(logs_dir.glob("*.csv")):
        if path.name.lower() == "list_of_prompt_dirs.csv":
            continue
        if path.name.endswith("_sim.csv"):
            continue
        files.append(path)
    return files


def calculate_total_size_mb(csv_path: Path) -> float:
    total_bytes = 0
    with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            value = row.get("Folder size in bytes", "0").strip()
            if not value:
                continue
            try:
                total_bytes += int(float(value))
            except ValueError:
                continue
    return round(total_bytes / (1024 ** 2), 2)


def ensure_workbook(output_file: Path):
    if output_file.exists():
        workbook = load_workbook(output_file)
    else:
        workbook = Workbook()

    sheet = workbook.active
    if sheet.max_row == 1 and not any(cell.value for cell in sheet[1]):
        sheet.title = "Stats"
        sheet.append(["Target folder name", "Total Size (MB) deleted"])
    elif sheet.max_row == 0:
        sheet.title = "Stats"
        sheet.append(["Target folder name", "Total Size (MB) deleted"])
    elif sheet.max_row > 0 and sheet.cell(row=1, column=1).value != "Target folder name":
        sheet.title = "Stats"
        sheet.insert_rows(1)
        sheet.append(["Target folder name", "Total Size (MB) deleted"])
    return workbook, sheet


def append_stats(output_file: Path, rows):
    workbook, sheet = ensure_workbook(output_file)
    for target_name, total_mb in rows:
        sheet.append([target_name, total_mb])
    workbook.save(output_file)


def main():
    csv_files = iter_log_csv_files(LOGS_DIR)
    if not csv_files:
        print(f"No eligible CSV files found in {LOGS_DIR}")
        return

    rows = []
    for csv_path in csv_files:
        total_mb = calculate_total_size_mb(csv_path)
        target_name = csv_path.stem
        rows.append((target_name, total_mb))

    append_stats(OUTPUT_FILE, rows)
    print(f"Wrote {len(rows)} rows to {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
